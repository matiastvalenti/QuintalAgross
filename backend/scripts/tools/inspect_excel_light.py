import openpyxl

file_path = 'excel/Stock.xlsx'
try:
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    print(f"Sheets: {wb.sheetnames}")
    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        print(f"\nSheet: {sheet_name}")
        rows = list(sheet.iter_rows(max_row=5, values_only=True))
        if rows:
            print(f"Headers (Row 1): {rows[0]}")
            print(f"Row 2: {rows[1] if len(rows) > 1 else 'N/A'}")
except Exception as e:
    print(f"Error: {e}")
