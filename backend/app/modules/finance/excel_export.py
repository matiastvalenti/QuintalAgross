from __future__ import annotations
from io import BytesIO
from datetime import datetime
from typing import List
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from app.db.models.finance_models import Cheque


def export_cheques_to_excel(cheques: List[Cheque]) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Cheques"
    
    header_font = Font(bold=True, color="FFFFFF", size=12)
    header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    
    border_side = Side(style="thin", color="CBD5E1")
    border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
    
    cell_alignment = Alignment(horizontal="left", vertical="center")
    
    headers = [
        "ID", "Banco", "Nro. Cheque", "Tipo", "CUIT Emisor", 
        "Cliente/Dador", "Beneficiario", "Importe", "Moneda",
        "F. Pago", "F. Movimiento", "F. Vencimiento", 
        "Movimiento", "Estado", "Rechazado", "ND Realizada",
        "Entregado a", "Fecha Entrega", "Nro. Orden Pago",
        "Creado", "Actualizado"
    ]
    
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = border
    
    for row_num, cheque in enumerate(cheques, 2):
        data = [
            cheque.id,
            cheque.banco,
            cheque.nro_cheque,
            cheque.tipo or "",
            cheque.cuit_emisor,
            cheque.cliente_dador or "",
            cheque.beneficiario or "",
            float(cheque.importe) if cheque.importe else 0,
            cheque.moneda or "",
            cheque.f_pago.strftime("%Y-%m-%d") if cheque.f_pago else "",
            cheque.f_movimiento.strftime("%Y-%m-%d") if cheque.f_movimiento else "",
            cheque.f_vencimiento.strftime("%Y-%m-%d") if cheque.f_vencimiento else "",
            cheque.movimiento or "",
            cheque.estado,
            "Sí" if cheque.rechazado else "No",
            "Sí" if cheque.nd_realizada else "No",
            cheque.entregado_a or "",
            cheque.fecha_entrega.strftime("%Y-%m-%d") if cheque.fecha_entrega else "",
            cheque.nro_orden_pago or "",
            cheque.created_at.strftime("%Y-%m-%d %H:%M") if cheque.created_at else "",
            cheque.updated_at.strftime("%Y-%m-%d %H:%M") if cheque.updated_at else "",
        ]
        
        for col_num, value in enumerate(data, 1):
            cell = ws.cell(row=row_num, column=col_num, value=value)
            cell.alignment = cell_alignment
            cell.border = border
            
            if col_num == 8 and value:
                cell.number_format = '$#,##0.00'
    
    column_widths = [8, 25, 15, 12, 18, 25, 25, 15, 10, 12, 15, 15, 15, 12, 10, 12, 20, 15, 18, 18, 18]
    for col_num, width in enumerate(column_widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=col_num).column_letter].width = width
    
    ws.freeze_panes = "A2"
    
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    return output
